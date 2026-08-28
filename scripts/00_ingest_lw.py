#!/usr/bin/env python3
"""Pull the NY Fed's Holston-Laubach-Williams (HLW) natural-rate-of-interest
estimate for the US -- a one-sided (real-time, filtered) r* series, updated
quarterly, from Holston, Laubach & Williams (2017/2023).

The workbook is organized as one sheet per historical *vintage* (real-time
estimates as they looked when first computed each quarter) -- we only want
the latest vintage's full history, which is the last quarter-labeled sheet
(e.g. "2026Q2"), not any of the earlier vintage snapshots.

Usage: python3 00_ingest_lw.py
"""
import io
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
import requests

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
LW_URL = (
    "https://www.newyorkfed.org/medialibrary/media/research/economists/"
    "williams/data/Holston_Laubach_Williams_real_time_estimates.xlsx"
)
VINTAGE_RE = re.compile(r"^\d{4}Q[1-4]$")


def main():
    resp = requests.get(LW_URL, timeout=60, headers={"User-Agent": "Mozilla/5.0"})
    resp.raise_for_status()
    wb = openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True)

    vintage_sheets = sorted(s for s in wb.sheetnames if VINTAGE_RE.match(s))
    if not vintage_sheets:
        sys.exit(f"No vintage-labeled sheets found; got {wb.sheetnames}")
    latest = vintage_sheets[-1]
    ws = wb[latest]

    rows = list(ws.iter_rows(values_only=True))
    header_row_idx = next(
        i for i, r in enumerate(rows) if r and r[0] == "Date"
    )
    header = rows[header_row_idx]
    # "Natural Rate (r*)" block starts a few columns in; the "US" sub-column
    # directly under it holds the one-sided US r* estimate.
    rstar_block_start = next(
        i for i, r in enumerate(rows[header_row_idx - 1]) if r == "Natural Rate (r*)"
    )
    us_col = rstar_block_start + header[rstar_block_start:].index("US")

    out_rows = []
    for r in rows[header_row_idx + 1:]:
        d = r[0]
        v = r[us_col]
        if d is None or v in (None, "NA"):
            continue
        out_rows.append((d.strftime("%Y-%m-%d"), float(v)))

    if not out_rows:
        sys.exit(f"Parsed 0 rows from sheet {latest!r} -- check column layout")

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    out_path = DATA_DIR / f"{date.today():%Y%m%d}_lw_rstar.csv"
    with open(out_path, "w") as f:
        f.write("date,lw_rstar\n")
        for d, v in out_rows:
            f.write(f"{d},{v}\n")
    print(f"Wrote {out_path} ({len(out_rows)} quarters, {out_rows[0][0]}..{out_rows[-1][0]}, "
          f"from vintage sheet {latest!r} of {len(vintage_sheets)} available)")


if __name__ == "__main__":
    main()
